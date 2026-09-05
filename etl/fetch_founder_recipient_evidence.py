#!/usr/bin/env python3
"""Fetch and normalize recipient-level founder school budget documents."""

from __future__ import annotations

import argparse
import csv
import hashlib
import json
import re
import unicodedata
from difflib import SequenceMatcher
from pathlib import Path
from typing import Any
from urllib.request import Request, urlopen

from openpyxl import load_workbook


ROOT = Path(__file__).resolve().parents[1]
RAW_ROOT = ROOT / "etl" / "data" / "raw"
CACHE_DIR = ROOT / "etl" / "data" / "founder_source_cache"
DEFAULT_CATALOG = ROOT / "etl" / "data" / "founder_source_catalog.csv"
DEFAULT_ALIASES = ROOT / "etl" / "data" / "founder_source_aliases.csv"
USER_AGENT = "cz-school-sankey/1.0"


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(
        description="Build school-level founder evidence from cataloged documents"
    )
    parser.add_argument("--year", action="append", type=int, required=True)
    parser.add_argument("--catalog", type=Path, default=DEFAULT_CATALOG)
    parser.add_argument("--aliases", type=Path, default=DEFAULT_ALIASES)
    parser.add_argument("--no-cache", action="store_true")
    return parser.parse_args()


def normalized_name(value: object) -> str:
    text = unicodedata.normalize("NFKD", str(value or "").casefold())
    text = "".join(char for char in text if not unicodedata.combining(char))
    text = text.replace("prispevkova organizace", "")
    text = re.sub(r"\b(praha 10|prahy 10)\b", "", text)
    text = re.sub(r"[^a-z0-9]+", " ", text)
    return " ".join(text.split())


def read_csv(path: Path) -> list[dict[str, str]]:
    with path.open(encoding="utf-8-sig", newline="") as handle:
        return list(csv.DictReader(handle))


def download_source(source: dict[str, str], no_cache: bool) -> Path:
    url = source["source_url"]
    suffix = f".{source.get('source_format') or 'bin'}"
    digest = hashlib.sha256(url.encode()).hexdigest()[:10]
    founder_ico = source["founder_id"].removeprefix("founder:")
    target = CACHE_DIR / f"{source['reporting_year']}-{founder_ico}-{digest}{suffix}"
    CACHE_DIR.mkdir(parents=True, exist_ok=True)
    if target.exists() and not no_cache:
        return target
    request = Request(url, headers={"User-Agent": USER_AGENT})
    with urlopen(request, timeout=90) as response:
        target.write_bytes(response.read())
    return target


def parse_school_budget_workbook(
    path: Path,
    *,
    year: int,
    unit_multiplier: int,
) -> list[dict[str, Any]]:
    workbook = load_workbook(path, data_only=True, read_only=True)
    if str(year) not in workbook.sheetnames:
        raise RuntimeError(f"Workbook {path.name} has no {year} sheet")
    worksheet = workbook[str(year)]

    header_row = name_column = amount_column = None
    for row_number, row in enumerate(worksheet.iter_rows(values_only=True), 1):
        normalized = [normalized_name(value) for value in row]
        for column, value in enumerate(normalized):
            if "nazev prispevkove organizace" in value:
                header_row = row_number
                name_column = column
            if "vynosy z rozpoctu mc" in value:
                amount_column = column
        if header_row is not None and name_column is not None and amount_column is not None:
            break
    if header_row is None or name_column is None or amount_column is None:
        raise RuntimeError(f"Could not find recipient and municipal-revenue columns in {path.name}")

    rows: list[dict[str, Any]] = []
    for values in worksheet.iter_rows(min_row=header_row + 1, values_only=True):
        recipient = values[name_column] if name_column < len(values) else None
        amount = values[amount_column] if amount_column < len(values) else None
        if not isinstance(recipient, str) or not isinstance(amount, (int, float)):
            continue
        rows.append(
            {
                "source_recipient_name": recipient.strip(),
                "amount": int(round(float(amount) * unit_multiplier)),
            }
        )
    return rows


def match_recipients(
    *,
    source_rows: list[dict[str, Any]],
    registry_rows: list[dict[str, str]],
    aliases: dict[str, str],
) -> tuple[list[dict[str, Any]], list[dict[str, Any]], list[dict[str, str]]]:
    registry_by_id = {row["institution_id"]: row for row in registry_rows}
    registry_by_name: dict[str, list[dict[str, str]]] = {}
    for row in registry_rows:
        registry_by_name.setdefault(normalized_name(row["institution_name"]), []).append(row)

    matched: list[dict[str, Any]] = []
    unmatched_source: list[dict[str, Any]] = []
    used_ids: set[str] = set()
    for source_row in source_rows:
        source_name = source_row["source_recipient_name"]
        normalized_source = normalized_name(source_name)
        match = None
        method = ""

        alias_id = aliases.get(normalized_source)
        if alias_id:
            match = registry_by_id.get(alias_id)
            method = "catalog_alias"
        elif len(registry_by_name.get(normalized_source, [])) == 1:
            match = registry_by_name[normalized_source][0]
            method = "normalized_exact"
        else:
            candidates = sorted(
                (
                    (SequenceMatcher(None, normalized_source, name).ratio(), row)
                    for name, rows in registry_by_name.items()
                    for row in rows
                    if row["institution_id"] not in used_ids
                ),
                key=lambda pair: pair[0],
                reverse=True,
            )
            if candidates:
                best_score, best = candidates[0]
                second_score = candidates[1][0] if len(candidates) > 1 else 0.0
                if best_score >= 0.85 and best_score - second_score >= 0.05:
                    match = best
                    method = "name_similarity"

        if match is None or match["institution_id"] in used_ids:
            unmatched_source.append(source_row)
            continue
        used_ids.add(match["institution_id"])
        matched.append({**source_row, **match, "match_method": method})

    unmatched_registry = [
        row for row in registry_rows if row["institution_id"] not in used_ids
    ]
    return matched, unmatched_source, unmatched_registry


def load_aliases(path: Path, year: int, founder_id: str) -> dict[str, str]:
    return {
        normalized_name(row["source_recipient_name"]): row["institution_id"]
        for row in read_csv(path)
        if int(row["reporting_year"]) == year and row["founder_id"] == founder_id
    }


def build_year(year: int, catalog: Path, aliases_path: Path, no_cache: bool) -> None:
    sources = [row for row in read_csv(catalog) if int(row["reporting_year"]) == year]
    if not sources:
        print(f"No recipient-level founder sources cataloged for {year}")
        return

    registry = read_csv(RAW_ROOT / str(year) / "school_entities.csv")
    evidence: list[dict[str, Any]] = []
    unmatched: list[dict[str, Any]] = []
    metadata_sources: list[dict[str, Any]] = []
    used_institutions: set[str] = set()

    for source in sources:
        founder_registry = [row for row in registry if row["founder_id"] == source["founder_id"]]
        archive = download_source(source, no_cache)
        if source["parser_profile"] != "school_budget_revenue_column":
            raise RuntimeError(f"Unsupported parser profile: {source['parser_profile']}")
        source_rows = parse_school_budget_workbook(
            archive,
            year=year,
            unit_multiplier=int(source["unit_multiplier"]),
        )
        matched, unmatched_source, unmatched_registry = match_recipients(
            source_rows=source_rows,
            registry_rows=founder_registry,
            aliases=load_aliases(aliases_path, year, source["founder_id"]),
        )
        if unmatched_registry:
            names = ", ".join(row["institution_name"] for row in unmatched_registry[:5])
            raise RuntimeError(
                f"{source['founder_id']} {year} leaves {len(unmatched_registry)} registry schools unmatched: {names}"
            )

        for row in matched:
            institution_id = row["institution_id"]
            if institution_id in used_institutions:
                raise RuntimeError(f"Multiple founder sources matched {institution_id} in {year}")
            used_institutions.add(institution_id)
            evidence.append(
                {
                    "institution_id": institution_id,
                    "founder_id": source["founder_id"],
                    "amount": row["amount"],
                    "basis": source["basis"],
                    "certainty": "observed",
                    "attribution_method": "recipient_reported_founder_budget",
                    "source_document_kind": source["document_kind"],
                    "source_recipient_name": row["source_recipient_name"],
                    "match_method": row["match_method"],
                    "source_url": source["source_url"],
                    "note": (
                        f"{source['founder_name']} approved {year} school budget; "
                        "recipient-level revenue from the founder budget"
                    ),
                }
            )
        unmatched.extend(
            {
                "reporting_year": year,
                "founder_id": source["founder_id"],
                "source_recipient_name": row["source_recipient_name"],
                "amount": row["amount"],
                "reason": "source_recipient_not_in_school_registry",
                "source_url": source["source_url"],
            }
            for row in unmatched_source
        )
        metadata_sources.append(
            {
                "founder_id": source["founder_id"],
                "source_url": source["source_url"],
                "source_row_count": len(source_rows),
                "matched_school_count": len(matched),
                "unmatched_source_count": len(unmatched_source),
            }
        )

    output_dir = RAW_ROOT / str(year)
    output_path = output_dir / "founder_recipient_evidence.csv"
    fields = [
        "institution_id",
        "founder_id",
        "amount",
        "basis",
        "certainty",
        "attribution_method",
        "source_document_kind",
        "source_recipient_name",
        "match_method",
        "source_url",
        "note",
    ]
    with output_path.open("w", encoding="utf-8", newline="") as handle:
        writer = csv.DictWriter(handle, fieldnames=fields, lineterminator="\n")
        writer.writeheader()
        writer.writerows(sorted(evidence, key=lambda row: row["institution_id"]))

    unmatched_path = output_dir / "founder_recipient_evidence_unmatched.csv"
    unmatched_fields = [
        "reporting_year",
        "founder_id",
        "source_recipient_name",
        "amount",
        "reason",
        "source_url",
    ]
    with unmatched_path.open("w", encoding="utf-8", newline="") as handle:
        writer = csv.DictWriter(handle, fieldnames=unmatched_fields, lineterminator="\n")
        writer.writeheader()
        writer.writerows(unmatched)

    metadata_path = output_dir / "founder_recipient_evidence.meta.json"
    with metadata_path.open("w", encoding="utf-8") as handle:
        json.dump(
            {
                "reporting_year": year,
                "evidence_row_count": len(evidence),
                "unmatched_source_count": len(unmatched),
                "sources": metadata_sources,
            },
            handle,
            ensure_ascii=False,
            indent=2,
        )
        handle.write("\n")

    print(
        f"Founder recipient evidence {year}: {len(evidence)} matched schools, "
        f"{len(unmatched)} source-only recipients"
    )
    print(f"Wrote {output_path.relative_to(ROOT)}")


def main() -> None:
    args = parse_args()
    for year in sorted(set(args.year)):
        build_year(year, args.catalog, args.aliases, args.no_cache)


if __name__ == "__main__":
    main()
